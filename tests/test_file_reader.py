"""Unit tests for file parsing, format support, and DoS limits."""

import io
import docx
import pytest

from utils.file_reader import (
    MAX_FILE_SIZE_BYTES,
    compute_file_signature,
    extract_text_from_file,
)


def test_compute_file_signature_deterministic():
    """Verify file signature is identical for the same bytes."""
    data = b"Arbitrary contract byte content"
    sig1 = compute_file_signature(data, "agreement.pdf")
    sig2 = compute_file_signature(data, "agreement.pdf")
    assert sig1 == sig2
    assert len(sig1) == 64


def test_compute_file_signature_different_files():
    """Verify different file contents produce distinct signatures."""
    sig1 = compute_file_signature(b"File A", "doc.txt")
    sig2 = compute_file_signature(b"File B", "doc.txt")
    assert sig1 != sig2


def test_compute_file_signature_type_error():
    """Verify TypeError on invalid argument types."""
    with pytest.raises(TypeError):
        compute_file_signature("not bytes", "file.txt")  # type: ignore


def test_extract_txt_valid():
    """Verify valid UTF-8 plaintext extraction."""
    content = "NON-DISCLOSURE AGREEMENT\nSection 1. Confidentiality."
    text, err = extract_text_from_file(content.encode("utf-8"), "nda.txt")
    assert err is None
    assert "NON-DISCLOSURE AGREEMENT" in text


def test_extract_txt_latin1_encoding():
    """Verify Latin-1 encoded text extraction."""
    raw = "Agreement with currency symbol: £100,000".encode("latin-1")
    text, err = extract_text_from_file(raw, "financial.txt")
    assert err is None
    assert "£100,000" in text


def test_extract_empty_file_returns_error():
    """Verify 0-byte file returns descriptive error message."""
    text, err = extract_text_from_file(b"", "empty.txt")
    assert text == ""
    assert "empty (0 bytes)" in err.lower()


def test_extract_file_size_limit_rejection():
    """Verify files exceeding 15MB are rejected to prevent DoS."""
    oversized = b"x" * (MAX_FILE_SIZE_BYTES + 1024)
    text, err = extract_text_from_file(oversized, "huge.txt")
    assert text == ""
    assert "exceeds the maximum allowed limit of 15 MB" in err


def test_extract_unsupported_extension():
    """Verify unsupported file extensions are rejected with clear guidance."""
    text, err = extract_text_from_file(b"data", "contract.exe")
    assert text == ""
    assert "unsupported file extension" in err.lower()


def test_extract_corrupted_docx():
    """Verify corrupted DOCX archive structure returns safe error instead of crashing."""
    corrupted_bytes = b"PK\x03\x04not a valid zip or docx file"
    text, err = extract_text_from_file(corrupted_bytes, "corrupt.docx")
    assert text == ""
    assert "corrupted docx" in err.lower() or "invalid" in err.lower()


def test_extract_corrupted_pdf():
    """Verify corrupted PDF bytes returns safe error message."""
    corrupted_pdf = b"%PDF-1.4 garbage binary stream that cannot parse"
    text, err = extract_text_from_file(corrupted_pdf, "corrupt.pdf")
    assert text == ""
    assert "pdf" in err.lower()


def test_extract_valid_docx():
    """Verify extraction from valid in-memory DOCX document."""
    doc = docx.Document()
    doc.add_heading("Master Services Agreement", level=1)
    doc.add_paragraph("This agreement governs cloud services provided by Vendor.")
    
    # Add a table
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Term"
    table.rows[0].cells[1].text = "12 Months"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    text, err = extract_text_from_file(docx_bytes, "services.docx")
    assert err is None
    assert "Master Services Agreement" in text
    assert "cloud services" in text
    assert "12 Months" in text
