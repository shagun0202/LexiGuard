"""Unit tests for Word (.docx) export generation."""

import io
import docx
import pytest

from utils.doc_exporter import export_to_docx


def test_export_to_docx_type_error():
    """Verify TypeError when document_name is not str."""
    with pytest.raises(TypeError):
        export_to_docx(123)  # type: ignore


def test_export_to_docx_empty_payload():
    """Verify document exports successfully even with minimal or empty payloads."""
    doc_bytes = export_to_docx("Minimal_Agreement.pdf")
    assert isinstance(doc_bytes, bytes)
    assert len(doc_bytes) > 1000

    # Verify readable docx structure
    parsed_doc = docx.Document(io.BytesIO(doc_bytes))
    full_text = "\n".join(p.text for p in parsed_doc.paragraphs)
    assert "LexiGuard Legal Intelligence Report" in full_text
    assert "IMPORTANT LEGAL DISCLAIMER" in full_text or "educational" in full_text.lower()


def test_export_to_docx_with_complete_analysis():
    """Verify document generates tables, checklists, and styled sections for complete data."""
    summary_data = {
        "summary": "This is an executive summary of the agreement.",
        "key_points": ["First takeaway", "Second takeaway"],
        "glossary": [{"term": "Indemnity", "definition": "Security against loss", "context": "Sec 4"}]
    }
    risk_data = {
        "risk_score": 45,
        "risk_level": "Medium",
        "score_rationale": "Moderate risk due to non-compete.",
        "inconsistencies": [{"description": "Notice window discrepancy", "conflicting_sections": "3 vs 9", "severity": "Medium"}],
        "clauses": [{
            "title": "Non-Compete",
            "category": "Termination",
            "risk_level": "Medium",
            "risk_explanation": "12 month restriction.",
            "quote": "Executive shall not compete for 12 months.",
            "counter_proposal": "Reduce to 6 months."
        }]
    }
    action_data = {
        "checklist": [{"task": "Review with attorney", "category": "Legal", "completed": False}],
        "deadlines": [{"task": "Notice of renewal", "responsible_party": "Client", "deadline": "60 days prior"}],
        "lawyer_questions": ["What is the enforceability of the non-compete?"]
    }

    doc_bytes = export_to_docx(
        document_name="Employment_Agreement.docx",
        summary_data=summary_data,
        risk_data=risk_data,
        action_data=action_data,
    )

    assert isinstance(doc_bytes, bytes)
    assert len(doc_bytes) > 5000

    parsed_doc = docx.Document(io.BytesIO(doc_bytes))
    full_text = "\n".join(p.text for p in parsed_doc.paragraphs)
    assert "1. Executive Summary & Overview" in full_text
    assert "2. Risk Assessment & Clause Breakdown" in full_text
    assert "3. Action Plan & Legal Consultation Guide" in full_text
    assert len(parsed_doc.tables) >= 3
