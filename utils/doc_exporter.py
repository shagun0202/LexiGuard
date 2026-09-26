"""Professional Word (.docx) report generation module for LexiGuard.

Transforms multi-tab analysis data into a clean, executive-ready Word document
with structured tables, callout blocks, and prominent educational disclaimers.
"""

import io
from datetime import datetime
from typing import Any, Dict, List, Optional

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


def _set_cell_background(cell: Any, fill_hex: str) -> None:
    """Set the background color of a Word table cell.

    Args:
        cell: docx table cell object.
        fill_hex: Hexadecimal color string without '#' (e.g. '0F172A').
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def _set_cell_margins(cell: Any, top: int = 120, bottom: int = 120, left: int = 150, right: int = 150) -> None:
    """Set internal cell padding (margins) in dxa (1/20th of a point).

    Args:
        cell: docx table cell object.
        top: Top padding.
        bottom: Bottom padding.
        left: Left padding.
        right: Right padding.
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def export_to_docx(
    document_name: str,
    summary_data: Optional[Dict[str, Any]] = None,
    risk_data: Optional[Dict[str, Any]] = None,
    action_data: Optional[Dict[str, Any]] = None,
) -> bytes:
    """Generate a formatted Word (.docx) report from LexiGuard analysis data.

    Args:
        document_name: Name of the analyzed legal document.
        summary_data: Output dictionary from simplify_service (optional).
        risk_data: Output dictionary from risk_service (optional).
        action_data: Output dictionary from action_service (optional).

    Returns:
        Bytes of the generated .docx file.

    Raises:
        TypeError: If document_name is not a str.
        RuntimeError: If document generation fails.
    """
    if not isinstance(document_name, str):
        raise TypeError(f"document_name must be str, got {type(document_name).__name__}")

    try:
        doc = docx.Document()

        # Set page margins
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Document Header
        title = doc.add_paragraph()
        title_run = title.add_run("LexiGuard Legal Intelligence Report")
        title_run.font.name = "Arial"
        title_run.font.size = Pt(22)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(15, 23, 42)  # Dark slate
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT

        subtitle = doc.add_paragraph()
        sub_run = subtitle.add_run(f"Automated Analysis & Risk Assessment: {document_name}")
        sub_run.font.name = "Arial"
        sub_run.font.size = Pt(12)
        sub_run.font.color.rgb = RGBColor(100, 116, 139)  # Slate grey

        date_p = doc.add_paragraph()
        date_run = date_p.add_run(f"Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}")
        date_run.font.name = "Arial"
        date_run.font.size = Pt(10)
        date_run.font.italic = True
        date_run.font.color.rgb = RGBColor(100, 116, 139)

        # Prominent Educational Disclaimer Box
        disc_table = doc.add_table(rows=1, cols=1)
        disc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        disc_cell = disc_table.cell(0, 0)
        _set_cell_background(disc_cell, "F1F5F9")
        _set_cell_margins(disc_cell, top=140, bottom=140, left=180, right=180)
        
        disc_p = disc_cell.paragraphs[0]
        disc_p.paragraph_format.space_before = Pt(2)
        disc_p.paragraph_format.space_after = Pt(2)
        d_title = disc_p.add_run("IMPORTANT LEGAL DISCLAIMER:\n")
        d_title.font.name = "Arial"
        d_title.font.bold = True
        d_title.font.size = Pt(9.5)
        d_title.font.color.rgb = RGBColor(185, 28, 28)  # Deep red

        d_text = disc_p.add_run(
            "This report is produced by LexiGuard for educational and informational purposes only. "
            "LexiGuard is an automated AI tool and NOT a law firm. This document DOES NOT constitute legal advice. "
            "Laws vary by jurisdiction; always consult a licensed attorney before executing any legal agreement."
        )
        d_text.font.name = "Arial"
        d_text.font.size = Pt(9)
        d_text.font.color.rgb = RGBColor(51, 65, 85)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 1. Executive Summary & Plain Language Overview
        if summary_data:
            _add_section_heading(doc, "1. Executive Summary & Overview")
            
            exec_p = doc.add_paragraph()
            exec_p.add_run(summary_data.get("summary", "No summary provided.")).font.size = Pt(10.5)

            key_points = summary_data.get("key_points", [])
            if key_points:
                doc.add_heading("Key Takeaways", level=2)
                for pt in key_points:
                    bp = doc.add_paragraph(style="List Bullet")
                    bp.add_run(pt).font.size = Pt(10)

            glossary = summary_data.get("glossary", [])
            if glossary:
                doc.add_heading("Plain-Language Legal Glossary", level=2)
                g_table = doc.add_table(rows=1, cols=3)
                g_table.style = "Table Grid"
                g_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                hdr_cells = g_table.rows[0].cells
                hdr_cells[0].text = "Term / Legalese"
                hdr_cells[1].text = "Plain English Meaning"
                hdr_cells[2].text = "Document Context"
                for cell in hdr_cells:
                    _set_cell_background(cell, "0F172A")
                    _set_cell_margins(cell)
                    for r in cell.paragraphs[0].runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(9.5)

                for item in glossary:
                    row_cells = g_table.add_row().cells
                    row_cells[0].text = item.get("term", "")
                    row_cells[1].text = item.get("definition", "")
                    row_cells[2].text = item.get("context", "")
                    for cell in row_cells:
                        _set_cell_margins(cell)
                        for r in cell.paragraphs[0].runs:
                            r.font.size = Pt(9)

            doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 2. Risk Assessment & Clause Breakdown
        if risk_data:
            _add_section_heading(doc, "2. Risk Assessment & Clause Breakdown")

            score = risk_data.get("risk_score", 0)
            level = risk_data.get("risk_level", "Medium")
            score_p = doc.add_paragraph()
            s_run = score_p.add_run(f"Overall Contract Risk Score: {score}/100 ({level} Risk)\n")
            s_run.font.bold = True
            s_run.font.size = Pt(12)
            score_p.add_run(risk_data.get("score_rationale", "")).font.size = Pt(10)

            # Inconsistencies & Contradictions
            inconsistencies = risk_data.get("inconsistencies", [])
            if inconsistencies:
                doc.add_heading("Detected Internal Contradictions & Ambiguities", level=2)
                for inc in inconsistencies:
                    inc_p = doc.add_paragraph(style="List Bullet")
                    inc_run = inc_p.add_run(inc.get("description", str(inc)))
                    inc_run.font.size = Pt(10)
                    inc_run.font.color.rgb = RGBColor(185, 28, 28)

            # Clause Cards
            clauses = risk_data.get("clauses", [])
            if clauses:
                doc.add_heading("Detailed Clause Analysis", level=2)
                c_table = doc.add_table(rows=1, cols=4)
                c_table.style = "Table Grid"
                c_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                c_hdr = c_table.rows[0].cells
                c_hdr[0].text = "Clause & Category"
                c_hdr[1].text = "Risk Level"
                c_hdr[2].text = "Verbatim Quote"
                c_hdr[3].text = "Recommended Counter-Position"
                for cell in c_hdr:
                    _set_cell_background(cell, "1E293B")
                    _set_cell_margins(cell)
                    for r in cell.paragraphs[0].runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(9.5)

                for clause in clauses:
                    row_cells = c_table.add_row().cells
                    row_cells[0].text = f"{clause.get('title', '')}\n[{clause.get('category', 'Other')}]"
                    row_cells[1].text = clause.get("risk_level", "Medium")
                    row_cells[2].text = f"\"{clause.get('quote', '')}\""
                    row_cells[3].text = clause.get("counter_proposal", "Standardize terms.")
                    for cell in row_cells:
                        _set_cell_margins(cell)
                        for r in cell.paragraphs[0].runs:
                            r.font.size = Pt(8.5)

            doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 3. Action Plan & Next Steps
        if action_data:
            _add_section_heading(doc, "3. Action Plan & Legal Consultation Guide")

            checklist = action_data.get("checklist", [])
            if checklist:
                doc.add_heading("Pre-Signing Verification Checklist", level=2)
                for item in checklist:
                    ch_p = doc.add_paragraph()
                    ch_p.add_run("☐  ").font.bold = True
                    task_text = item.get("task", str(item)) if isinstance(item, dict) else str(item)
                    ch_p.add_run(task_text).font.size = Pt(10)

            deadlines = action_data.get("deadlines", [])
            if deadlines:
                doc.add_heading("Key Obligations & Milestones", level=2)
                d_table = doc.add_table(rows=1, cols=3)
                d_table.style = "Table Grid"
                d_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                d_hdr = d_table.rows[0].cells
                d_hdr[0].text = "Obligation / Milestone"
                d_hdr[1].text = "Responsible Party"
                d_hdr[2].text = "Timeline / Trigger"
                for cell in d_hdr:
                    _set_cell_background(cell, "0F172A")
                    _set_cell_margins(cell)
                    for r in cell.paragraphs[0].runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(9.5)

                for dl in deadlines:
                    row_cells = d_table.add_row().cells
                    row_cells[0].text = dl.get("task", "")
                    row_cells[1].text = dl.get("responsible_party", "")
                    row_cells[2].text = dl.get("deadline", "")
                    for cell in row_cells:
                        _set_cell_margins(cell)
                        for r in cell.paragraphs[0].runs:
                            r.font.size = Pt(9)

            questions = action_data.get("lawyer_questions", [])
            if questions:
                doc.add_heading("Targeted Questions for Legal Counsel", level=2)
                for idx, q in enumerate(questions, start=1):
                    q_p = doc.add_paragraph()
                    q_p.add_run(f"{idx}. ").font.bold = True
                    q_p.add_run(q).font.size = Pt(10)

        # Final Footer Note
        doc.add_paragraph().paragraph_format.space_after = Pt(20)
        footer_p = doc.add_paragraph()
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = footer_p.add_run("— End of LexiGuard Intelligence Report —\nStrictly Confidential & Educational")
        f_run.font.size = Pt(9)
        f_run.font.italic = True
        f_run.font.color.rgb = RGBColor(148, 163, 184)

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    except Exception as exc:
        raise RuntimeError(f"Failed to generate Word report: {str(exc)}") from exc


def _add_section_heading(doc: Any, text: str) -> None:
    """Helper to add a standardized styled section heading.

    Args:
        doc: docx Document object.
        text: Heading text.
    """
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(6)
    run = heading.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
